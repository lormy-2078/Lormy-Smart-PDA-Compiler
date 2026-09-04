"""
AI-powered field extraction using the Claude API -- native vision and PDF
understanding instead of the offline heuristic label-scanner in extraction.py.
This is the "real OCR / computer vision" path: it reads scanned/photographed
documents that have no text layer, which extraction.py cannot do at all.

Three entry points, all returning (found: dict, confidence: dict):
  * extract_vessel_particulars_ai(file_path)      -- one file (pdf/image/docx/xls)
  * extract_vessel_particulars_ai_from_text(text) -- pasted appointment email text
  * extract_vessel_particulars_ai_from_content(blocks) -- pre-built content blocks
        (used for emails: body text + each attachment as its own block)

Requires ANTHROPIC_API_KEY in the environment (see README). Callers should
catch exceptions here and fall back to extraction.py's heuristic scanner --
this module raises rather than silently degrading, so the fallback decision
stays visible to the caller (and the user, via a flash message).
"""

import base64
import json
import os
from pathlib import Path

import anthropic

from pda_engine.rules_engine import LINER_TERMS_OPTIONS

MODEL = os.environ.get("ANTHROPIC_MODEL", "claude-opus-4-8")

_TEXT_FIELDS = [
    "vessel", "imo", "flag", "vessel_type", "eta", "port_of_call", "cargo",
    "our_ref", "your_ref", "principal", "principal_attn", "receiver",
    "receiver_attn", "liner_terms", "appointment_text",
]
_NUMERIC_FIELDS = ["grt", "loa", "draft", "cargo_qty_mt", "port_stay_days"]

# Plain (non-union) types only -- the API caps schemas at 16 union/nullable-typed
# parameters (exponential compilation cost). Using "" / 0 as a not-found sentinel
# instead of nullable types avoids the limit entirely.
FIELD_SCHEMA = {
    "type": "object",
    "properties": {
        **{f: {"type": "string"} for f in _TEXT_FIELDS},
        **{f: {"type": "number"} for f in _NUMERIC_FIELDS},
        "liner_terms": {
            "type": "string",
            "description": (
                "The load/discharge (liner) terms if stated (e.g. Free Out, FIOS, LIFO, "
                "liner terms). Map to EXACTLY one of these values: "
                + "; ".join(LINER_TERMS_OPTIONS)
                + ". Empty string if the terms are not stated."
            ),
        },
        "appointment_text": {
            "type": "string",
            "description": (
                "Verbatim or closely paraphrased sentences bearing on billing / "
                "cost-allocation (FIO, LIFO, FILO, liner terms, port charges owner, "
                "receiver account, 'no separate cargo billing', etc). Empty string if none found."
            ),
        },
        "principal": {
            "type": "string",
            "description": (
                "The Principal's full BILLING address, not just a company name. Look for cues "
                "like 'Invoice to:', 'Bill to:', 'Please invoice:', or a letterhead/address block "
                "for the party being appointed/nominated. If the billing address includes a VAT/"
                "registration number, street, postal code, city, and/or country, capture ALL of "
                "them, each on its own line separated by a newline character (company name first "
                "line, then VAT/reg no., then street, then postal code + city, then country). If "
                "only a company name is stated with no further address detail, return just that name."
            ),
        },
        "eta": {"type": "string", "description": "Date as written in the document, e.g. 21-Jul-2026"},
    },
    "required": _TEXT_FIELDS + _NUMERIC_FIELDS,
    "additionalProperties": False,
}

EXTRACTION_PROMPT = (
    "This is one or more port-agency documents / emails for a vessel call (an appointment "
    "or nomination email, vessel particulars sheet, cargo details, or similar) -- text may "
    "be pasted email content, and any attachments may be scans or photos of printed pages. "
    "Read everything provided and extract the fields in the schema exactly as they appear. "
    "If several sources disagree, prefer the most specific/official one. Use an empty string "
    "\"\" for any text field not present or illegible, and 0 for any numeric field not present "
    "-- do not guess or invent a value. Numeric fields (grt, loa, draft, cargo_qty_mt, "
    "port_stay_days) must be plain numbers with no units, currency symbols, or thousands separators."
)

_IMAGE_MEDIA_TYPES = {
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
}


def is_configured() -> bool:
    return bool(os.environ.get("ANTHROPIC_API_KEY"))


def _docx_text(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    text = "\n".join(p.text for p in d.paragraphs)
    for table in d.tables:
        for row in table.rows:
            text += "\n" + " | ".join(c.text for c in row.cells)
    return text


def _xls_text(path: Path) -> str:
    """Flat cell dump of a spreadsheet, so Claude can read particulars sent as .xls/.xlsx."""
    rows = []
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        for ws in wb.worksheets:
            rows.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c not in (None, "")]
                if cells:
                    rows.append(" | ".join(cells))
    else:
        import xlrd
        wb = xlrd.open_workbook(str(path))
        for ws in wb.sheets():
            rows.append(f"[Sheet: {ws.name}]")
            for r in range(ws.nrows):
                cells = [str(ws.cell_value(r, c)) for c in range(ws.ncols) if ws.cell_value(r, c) not in (None, "")]
                if cells:
                    rows.append(" | ".join(cells))
    return "\n".join(rows)


def content_block_any(path):
    """A single Claude content block for any supported file type: PDF and images
    go in natively (real vision/OCR); docx and xls/xlsx are converted to a text
    block. Returns None for an unsupported type."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        data = base64.standard_b64encode(path.read_bytes()).decode("utf-8")
        return {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": data}}
    if ext in _IMAGE_MEDIA_TYPES:
        data = base64.standard_b64encode(path.read_bytes()).decode("utf-8")
        return {"type": "image", "source": {"type": "base64", "media_type": _IMAGE_MEDIA_TYPES[ext], "data": data}}
    if ext == ".docx":
        return {"type": "text", "text": f"[Attachment {path.name}]\n{_docx_text(path)}"}
    if ext in (".xls", ".xlsx"):
        return {"type": "text", "text": f"[Attachment {path.name}]\n{_xls_text(path)}"}
    return None


def _run_extraction(content_blocks, schema=FIELD_SCHEMA, prompt=EXTRACTION_PROMPT, numeric_fields=None, max_tokens=2048):
    """Core API call: content_blocks is a list of Claude content blocks (text /
    document / image). Appends the extraction instruction and returns
    (found, confidence). `schema`/`prompt`/`numeric_fields` default to the
    vessel-particulars extraction; other extractors (e.g. the BL/manifest one)
    pass their own."""
    numeric_fields = _NUMERIC_FIELDS if numeric_fields is None else numeric_fields
    client = anthropic.Anthropic()  # reads ANTHROPIC_API_KEY from the environment
    content = list(content_blocks) + [{"type": "text", "text": prompt}]

    response = client.messages.create(
        model=MODEL,
        max_tokens=max_tokens,
        output_config={"format": {"type": "json_schema", "schema": schema}},
        messages=[{"role": "user", "content": content}],
    )

    text = next(b.text for b in response.content if b.type == "text")
    data = json.loads(text)

    # "" / 0 are the not-found sentinels used in the schema (see FIELD_SCHEMA comment).
    found = {
        k: v for k, v in data.items()
        if not (v is None or v == "" or (k in numeric_fields and v == 0))
    }
    confidence = {k: "high" for k in found}
    return found, confidence


def extract_vessel_particulars_ai(file_path: str):
    """Extract from a single file (pdf/image/docx/xls/xlsx). Raises on missing
    API key, API error, or an unsupported file type -- callers should catch and
    fall back to extraction.extract_vessel_particulars()."""
    block = content_block_any(file_path)
    if block is None:
        raise ValueError(f"Unsupported file type for AI extraction: {Path(file_path).suffix}")
    return _run_extraction([block])


def extract_vessel_particulars_ai_from_text(text: str):
    """Extract from pasted appointment / nomination email text."""
    if not (text or "").strip():
        raise ValueError("No text provided for AI extraction.")
    return _run_extraction([{"type": "text", "text": f"Pasted appointment / email text:\n\n{text}"}])


def extract_vessel_particulars_ai_from_content(blocks):
    """Extract from pre-built content blocks (email body + attachments)."""
    if not blocks:
        raise ValueError("No content provided for AI extraction.")
    return _run_extraction(blocks)


# --- Cargo manifest / Bill of Lading extraction -----------------------------
# A second, later extraction stage (see README) -- narrower than the vessel-
# particulars schema above, focused on exactly what a manifest/BL states that
# the appointment email doesn't. A single cargo manifest routinely carries
# SEVERAL Bills of Lading (each an "MR number" line with its own receiver,
# weight and piece count) -- e.g. the real BALTIC ERICA manifest has 4 (GH005,
# BF003, BF004, GH006). So the schema returns a LIST of bills, one per line
# item, each becoming its own BL-scoped Receiver PDA.

BL_MANIFEST_SCHEMA = {
    "type": "object",
    "properties": {
        "bills": {
            "type": "array",
            "description": "One entry per Bill of Lading / MR number line on the manifest.",
            "items": {
                "type": "object",
                "properties": {
                    "bl_number": {
                        "type": "string",
                        "description": "The Bill of Lading number, MR number, or MR No. for this line. Empty string if none.",
                    },
                    "receiver": {
                        "type": "string",
                        "description": (
                            "The receiver/consignee to be billed the local charges -- the local (Ghana) "
                            "party, i.e. the Consignee, its full name and address. Empty string if none."
                        ),
                    },
                    "receiver_attn": {"type": "string", "description": "Attention/contact person, if named. Empty string if none."},
                    "pieces": {"type": "number", "description": "Number of pieces/bags on this line. 0 if not stated."},
                    "gross_weight_mt": {
                        "type": "number",
                        "description": "GROSS weight of this line in metric tonnes (convert from kg if stated in kg -- divide by 1000). 0 if not stated.",
                    },
                    "net_weight_mt": {
                        "type": "number",
                        "description": "NET weight of this line in metric tonnes (convert from kg if stated in kg). 0 if not stated.",
                    },
                    "bag_size_kg": {
                        "type": "number",
                        "description": (
                            "The stated size/weight of ONE bag in kg if the description gives it (e.g. "
                            "'BIG BAGS OF 1250 KG' -> 1250, '1000 KG' -> 1000). 0 if not stated."
                        ),
                    },
                },
                "required": ["bl_number", "receiver", "receiver_attn", "pieces", "gross_weight_mt", "net_weight_mt", "bag_size_kg"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["bills"],
    "additionalProperties": False,
}

BL_MANIFEST_PROMPT = (
    "This is a cargo manifest or Bill of Lading document. It may list SEVERAL Bills of Lading -- each "
    "'MR number' / 'MR No.' / BL number line is a separate bill with its own receiver, weight and piece "
    "count. Return one entry in `bills` per such line. For each: the BL/MR number, the receiver/consignee "
    "to be billed locally (the Consignee company name and address), the attention person if named, the "
    "number of pieces/bags, the GROSS weight in metric tonnes, the NET weight in metric tonnes (convert "
    "from kg where stated), and the stated bag size in kg if given in the goods description. Use \"\" for "
    "missing text and 0 for missing numbers -- do not guess. Do not include any manifest 'Total' row as a bill."
)


def _default_cbm_multiplier(weight_per_bag_mt: float) -> float:
    """DEFAULT GPHA CBM multiplier for a bag of a given weight -- 1.3 for bags
    close to 1.25mt/1250kg, 1.0 for bags close to 1mt/1000kg. Always meant to be
    user-edited per shipment; this is only a starting point."""
    if not weight_per_bag_mt:
        return 0.0
    return 1.3 if abs(weight_per_bag_mt - 1.25) <= abs(weight_per_bag_mt - 1.0) else 1.0


def derive_bl_fields(bill: dict) -> dict:
    """Turns one raw extracted manifest bill into the fields the Add-BL form
    needs. Charges are on GROSS weight (cargo_qty_mt = gross). weight-per-bag
    comes from the stated bag size if given, else net_weight/pieces -- and only
    picks the DEFAULT multiplier; CBM = pieces x multiplier (all editable)."""
    pieces = bill.get("pieces") or 0
    gross = bill.get("gross_weight_mt") or 0
    net = bill.get("net_weight_mt") or 0
    bag_size_kg = bill.get("bag_size_kg") or 0

    if bag_size_kg:
        weight_per_bag = bag_size_kg / 1000.0
    elif pieces:
        weight_per_bag = net / pieces
    else:
        weight_per_bag = 0.0
    multiplier = _default_cbm_multiplier(weight_per_bag)
    cbm = round(pieces * multiplier, 2)

    return {
        "bl_number": bill.get("bl_number", ""),
        "receiver": bill.get("receiver", ""),
        "receiver_attn": bill.get("receiver_attn", ""),
        "pieces": pieces,
        "cargo_qty_mt": gross,          # charges are on GROSS weight
        "net_weight_mt": net,
        "bag_size_kg": bag_size_kg,
        "cbm_multiplier": multiplier,
        "cargo_qty_cbm": cbm,
    }


def extract_bl_manifest_ai(file_path: str):
    """Extract every Bill of Lading on a cargo manifest / BL document
    (pdf/image/docx/xls/xlsx) -- returns a list of derived bill dicts (see
    derive_bl_fields), one per MR/BL line. Raises on missing API key, API
    error, or an unsupported file type -- callers should catch and fall back
    to manual entry (there is no offline heuristic for this one)."""
    block = content_block_any(file_path)
    if block is None:
        raise ValueError(f"Unsupported file type for BL/manifest extraction: {Path(file_path).suffix}")
    found, _conf = _run_extraction([block], schema=BL_MANIFEST_SCHEMA, prompt=BL_MANIFEST_PROMPT,
                                   numeric_fields=[], max_tokens=4096)
    bills = found.get("bills") or []
    return [derive_bl_fields(b) for b in bills]
