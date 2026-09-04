"""
Best-effort field extraction for uploaded vessel-particulars / appointment
documents: .xls / .xlsx spreadsheets, .pdf, and .docx. This is a heuristic
label-scan, not a real document-AI model (the PRD specifies Azure AI
Document Intelligence for that; this repo has no Azure credentials). Every
extracted field is returned with a confidence flag so the Review screen can
show the operator what to check.

Strategy: reduce every document to a stream of "rows", each a list of
text tokens in reading order (spreadsheet cells; PDF/Word table cells;
colon- or wide-gap-split segments of a text line). Then walk each row left
to right -- whenever a token matches a known label (e.g. "vessel", "imo",
"grt"), take the next token in that row as the value. Works for the
label/value layout used in the the company PDA templates (e.g.
TAK26063 KASSIOPI.GR.xls) and for "Label: Value" style appointment letters.

Scanned/image-only PDFs have no text layer and will not extract anything
here -- that needs OCR (e.g. Tesseract), which isn't installed in this
environment.
"""

import re
from pathlib import Path

LABELS = {
    "vessel": ["vessel"],
    "imo": ["imo"],
    "flag": ["flag"],
    "grt": ["grt"],
    "loa": ["loa"],
    "draft": ["draft"],
    "cargo": ["cargo"],
    "vessel_type": ["vessel type"],
    "eta": ["eta"],
    "port_of_call": ["port of call"],
    "cargo_qty_mt": ["cargo quantit", "provisional cargo quantit"],
    "port_stay_days": ["port stay", "provisional port stay"],
    "our_ref": ["our ref"],
    "your_ref": ["your ref"],
    "principal": ["to:"],
    "principal_attn": ["att:"],
    "agency_fee": ["agency fee"],
}

# Fields whose label must match the whole cell text exactly (after stripping
# a trailing colon), because their pattern is a short substring of several
# other labels (e.g. "cargo" inside "Actual Cargo Quantity").
EXACT_MATCH_FIELDS = {"cargo"}


def _rows_from_xls(path):
    import xlrd
    wb = xlrd.open_workbook(path)
    for sheet in wb.sheets():
        for r in range(sheet.nrows):
            yield [sheet.cell_value(r, c) for c in range(sheet.ncols)]


def _rows_from_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            yield list(row)


def _lines_to_rows(lines):
    """Turn free-text lines into label/value token rows.

    Handles "Label: Value" and wide-gap "Label    Value" layouts on one
    line. If a line is just a bare label (colon, nothing after it), the
    next non-empty line is pulled in as its value -- common in appointment
    letters where the value sits on its own line below the heading.
    """
    rows = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if ":" in line:
            label, _, rest = line.partition(":")
            label, rest = label.strip(), rest.strip()
            if rest:
                rows.append([label, rest])
                continue
            # bare "Label:" -- look ahead to the next non-empty line for the value,
            # but only if that line isn't itself a "Label: Value" line of its own.
            j = i
            while j < n and not lines[j].strip():
                j += 1
            if j < n and ":" not in lines[j]:
                rows.append([label, lines[j].strip()])
                i = j + 1
            else:
                rows.append([label])
            continue
        parts = [p.strip() for p in re.split(r"\s{2,}|\t", line) if p.strip()]
        rows.append(parts if len(parts) >= 2 else [line])
    return rows


def _rows_from_pdf(path):
    import pdfplumber
    rows = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in (page.extract_tables() or []):
                for r in table:
                    rows.append([c if c is not None else "" for c in r])
            text = page.extract_text() or ""
            rows.extend(_lines_to_rows(text.splitlines()))
    return rows


def _rows_from_docx(path):
    import docx
    d = docx.Document(path)
    rows = []
    for table in d.tables:
        for r in table.rows:
            rows.append([c.text for c in r.cells])
    lines = [p.text for p in d.paragraphs]
    rows.extend(_lines_to_rows(lines))
    return rows


def _clean(v):
    if isinstance(v, str):
        return v.strip()
    return v


def extract_vessel_particulars(file_path: str):
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext == ".xls":
        row_iter = _rows_from_xls(str(path))
    elif ext in (".xlsx", ".xlsm"):
        row_iter = _rows_from_xlsx(str(path))
    elif ext == ".pdf":
        row_iter = _rows_from_pdf(str(path))
    elif ext == ".docx":
        row_iter = _rows_from_docx(str(path))
    else:
        return {}, {}

    found = {}
    confidence = {}

    def set_field(field, candidate):
        if field in found:
            return
        found[field] = candidate
        confidence[field] = "high" if candidate not in (None, "") else "low"

    for row in row_iter:
        cells = [(i, _clean(v)) for i, v in enumerate(row) if _clean(v) not in (None, "")]
        for idx, (col, val) in enumerate(cells):
            if not isinstance(val, str):
                continue
            low = val.lower()

            # Combined "ETA / Flag" style header: two values follow in the next two cells.
            if "eta" in low and "flag" in low:
                if idx + 1 < len(cells):
                    set_field("eta", cells[idx + 1][1])
                if idx + 2 < len(cells):
                    set_field("flag", cells[idx + 2][1])
                continue

            # Pick only the longest (most specific) matching label for this cell,
            # so e.g. "Provisional Cargo Quantitiy" maps to cargo_qty_mt, not the
            # generic "cargo" (=cargo type) field.
            stripped = low.strip().rstrip(":").strip()
            candidates = []
            for field, patterns in LABELS.items():
                if field in found:
                    continue
                for p in patterns:
                    if field in EXACT_MATCH_FIELDS:
                        if stripped == p:
                            candidates.append((len(p), field))
                    elif p in low:
                        candidates.append((len(p), field))
            if candidates:
                candidates.sort(reverse=True)
                _, field = candidates[0]
                if idx + 1 < len(cells):
                    set_field(field, cells[idx + 1][1])

    # light normalization
    if "imo" in found:
        m = re.search(r"\d{6,7}", str(found["imo"]))
        if m:
            found["imo"] = m.group(0)
    for numeric_field in ("grt", "loa", "draft", "cargo_qty_mt"):
        if numeric_field in found:
            m = re.search(r"[\d,.]+", str(found[numeric_field]))
            if m:
                try:
                    found[numeric_field] = float(m.group(0).replace(",", ""))
                except ValueError:
                    pass
    if "port_stay_days" in found:
        m = re.search(r"[\d.]+", str(found["port_stay_days"]))
        if m:
            try:
                found["port_stay_days"] = float(m.group(0))
            except ValueError:
                pass

    if "eta" in found:
        val = found["eta"]
        if hasattr(val, "strftime"):
            found["eta"] = val.strftime("%d-%b-%Y")
        elif isinstance(val, (int, float)):
            import datetime
            found["eta"] = (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(val))).strftime("%d-%b-%Y")

    return found, confidence
